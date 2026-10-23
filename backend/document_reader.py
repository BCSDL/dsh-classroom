"""Shared CPU document extraction; no classroom runtime or model imports."""
import json
import os
import subprocess
import sys
from pathlib import Path

def document(path, output_dir):
    """Page-sized output; preserve source names, pages, and images for visual analysis."""
    path = Path(path).resolve()
    if not path.is_file():
        raise FileNotFoundError(str(path))
    suffix = path.suffix.lower()
    if suffix in {'.dxf', '.step', '.stp', '.iges', '.igs', '.stl', '.obj', '.ply', '.off', '.glb', '.gltf'}:
        # Native CAD libraries print to C stdout. Isolate them so MCP/worker
        # protocol lines stay clean even while another job is running.
        process = subprocess.run([sys.executable, str(Path(__file__).with_name('engineering.py')),
                                  str(path), str(output_dir)], capture_output=True,
                                 creationflags=subprocess.CREATE_NO_WINDOW if os.name == 'nt' else 0)
        parsed = None
        for line in process.stdout.decode('utf-8', errors='replace').splitlines():
            try:
                value = json.loads(line)
                if isinstance(value, dict) and ('result' in value or 'error' in value):
                    parsed = value
            except ValueError:
                pass
        if not parsed or process.returncode or 'error' in parsed:
            raise RuntimeError((parsed or {}).get('error', '工程解析失败，请检查文件和 scripts/setup.ps1 -Engineering'))
        yield from parsed['result']
    elif suffix in {'.dwg', '.sldprt', '.sldasm', '.catpart', '.catproduct', '.prt', '.pcbdoc', '.schdoc'}:
        raise ValueError('专有工业格式需要对应厂商软件或授权转换器，请导出 STEP/IGES/DXF/PDF；不会假装读取成功')
    elif suffix == '.pdf':
        import fitz
        with fitz.open(path) as pdf:
            for index, page in enumerate(pdf):
                image = output_dir / f'{path.stem}-{index + 1}.jpg'
                scale = min(1.5, 1920 / max(page.rect.width, page.rect.height))
                page.get_pixmap(matrix=fitz.Matrix(scale, scale)).save(image)
                yield {'source': path.name, 'page': index + 1, 'text': page.get_text(), 'image': str(image)}
    elif suffix in {'.png', '.jpg', '.jpeg', '.webp', '.bmp', '.tif', '.tiff'}:
        from PIL import Image
        image = output_dir / (path.stem + '.jpg')
        with Image.open(path) as picture:
            picture.thumbnail((1920, 1920))
            picture.convert('RGB').save(image)
        yield {'source': path.name, 'page': 1, 'text': '', 'image': str(image)}
    elif suffix == '.pptx':
        from pptx import Presentation
        deck = Presentation(path)
        for number, slide in enumerate(deck.slides, 1):
            text = '\n'.join(shape.text for shape in slide.shapes if shape.has_text_frame)
            if slide.has_notes_slide:
                text += '\n' + slide.notes_slide.notes_text_frame.text
            yield {'source': path.name, 'page': number, 'text': text,
                   'coverage': 'slide text and notes; export to PDF for full visual layout'}
            for part, shape in enumerate(slide.shapes):
                if hasattr(shape, 'image'):
                    image = output_dir / f'{path.stem}-{number}-{part}.{shape.image.ext}'
                    image.write_bytes(shape.image.blob)
                    yield {'source': path.name, 'page': number, 'text': '', 'image': str(image)}
    elif suffix == '.docx':
        from docx import Document
        doc = Document(path)
        texts = [p.text for p in doc.paragraphs] + [' | '.join(c.text for c in row.cells) for table in doc.tables for row in table.rows]
        for index in range(0, len(texts), 100):
            yield {'source': path.name, 'block': index // 100 + 1, 'text': '\n'.join(texts[index:index + 100])}
    elif suffix in {'.xlsx', '.xlsm'}:
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            for sheet in workbook:
                rows = []
                for number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                    rows.append(f"{number}: " + ' | '.join('' if cell is None else str(cell) for cell in row))
                    if len(rows) == 100:
                        yield {'source': path.name, 'sheet': sheet.title, 'text': '\n'.join(rows)}
                        rows.clear()
                if rows:
                    yield {'source': path.name, 'sheet': sheet.title, 'text': '\n'.join(rows)}
        finally:
            workbook.close()
    else:
        # Reject opaque/binary formats rather than pretending an extension implies comprehension.
        with path.open('rb') as handle:
            sample = handle.read(8192)
        if sample.startswith((b'\xff\xfe', b'\xfe\xff')):
            encoding = 'utf-16'
        elif b'\0' in sample:
            raise ValueError(f'{suffix} 是二进制格式，需要该格式的专用转换器')
        else:
            encoding = 'utf-8-sig'
            try:
                sample.decode(encoding)
            except UnicodeError:
                encoding = 'gb18030'
                sample.decode(encoding)
        with path.open(encoding=encoding) as handle:
            block = []
            for line in handle:
                block.append(line)
                if sum(len(x) for x in block) >= 18000:
                    yield {'source': path.name, 'text': ''.join(block)}
                    block.clear()
            if block:
                yield {'source': path.name, 'text': ''.join(block)}

